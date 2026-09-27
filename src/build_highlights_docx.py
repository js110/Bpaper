"""Create the Elsevier Highlights Word file from paper/highlights.txt."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

source=Path('paper/highlights.txt')
lines=[x.strip().lstrip('•').strip() for x in source.read_text().splitlines()
       if x.strip() and x.strip().lower()!='highlights']
if not 3<=len(lines)<=5:raise ValueError('Elsevier highlights require 3-5 bullets')
if any(len(x)>85 for x in lines):raise ValueError('highlight exceeds 85 characters')
doc=Document()
style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(11)
doc.add_heading('Highlights',level=1)
for line in lines:
    doc.add_paragraph(line,style='List Bullet')
doc.save('paper/highlights.docx')
print({'bullets':len(lines),'max_chars':max(map(len,lines))})
